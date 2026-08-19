from io import BytesIO

from docx import Document

from app.rag.extractors.base_extractor import BaseExtractor
from app.rag.models import ExtractedDocument, ExtractedSegment


class DocxExtractor(BaseExtractor):
    """
    Bộ trích xuất văn bản từ tệp Word (.docx) bảo toàn cấu trúc phân cấp (Hierarchical Heading Preservation).

    Cơ chế hoạt động và giữ ngữ cảnh cây đề mục (Heading Breadcrumb):
    - Đọc tệp Word trực tiếp từ bộ nhớ đệm byte (`BytesIO`) bằng thư viện `python-docx`.
    - Quét tuần tự qua từng paragraph và phân tích `style.name`.
    - Khi gặp các đề mục (Heading 1, Heading 2, Heading 3...), bộ trích xuất cập nhật danh sách `headings` theo đúng cấp độ (Level).
    - Mỗi đoạn văn bản nội dung sau đó được gắn kèm đường dẫn đề mục phân cấp trong metadata (`headingPath`, ví dụ: "Chương 1 > Phần A > Mục 1.1"),
      giúp mô hình RAG hiểu rõ ngữ cảnh phạm vi của đoạn văn bản đó dù tài liệu rất dài.
    """

    async def extract(
        self, *, content: str | None = None, file_bytes: bytes | None = None
    ) -> ExtractedDocument:
        """
        Phân tích tệp DOCX và chuyển đổi thành danh sách các ExtractedSegment có metadata headingPath.

        Cơ chế:
        1. Kiểm tra tính hợp lệ của `file_bytes`.
        2. Duyệt qua tất cả các đoạn văn bản trong tệp docx.
        3. Cập nhật cây tiêu đề khi gặp style dạng Heading.
        4. Tạo `ExtractedSegment` cho các đoạn văn bản thông thường.

        Args:
            content (str | None): Không sử dụng cho tệp DOCX.
            file_bytes (bytes | None): Dữ liệu nhị phân của tệp .docx.

        Returns:
            ExtractedDocument: Tài liệu chứa các segment kèm `headingPath`.

        Raises:
            ValueError: Nếu thiếu `file_bytes` hoặc tệp DOCX không chứa nội dung văn bản.
        """
        if not file_bytes:
            raise ValueError("DOCX phải có dữ liệu file")
        headings: list[str] = []
        segments: list[ExtractedSegment] = []
        document = Document(BytesIO(file_bytes))
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except ValueError:
                    level = 1
                headings = headings[: level - 1]
                headings.append(text)
                continue
            segments.append(
                ExtractedSegment(
                    text=text, metadata={"headingPath": " > ".join(headings)}
                )
            )
        if not segments:
            raise ValueError("DOCX không chứa nội dung văn bản")
        return ExtractedDocument(segments=segments)
